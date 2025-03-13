#!/usr/bin/env python3
import os
from pptx import Presentation
from google.cloud import translate_v2 as translate
from dotenv import load_dotenv  # Import load_dotenv

# Load environment variables from .env file
load_dotenv()

def translate_text(client, text, target='en', source='no'):
    result = client.translate(text, target_language=target, source_language=source)
    return result['translatedText']

def main():
    file_path = input("Enter the full path to the .pptx file: ").strip()
    if not os.path.isfile(file_path):
        print("The file does not exist. Please check the path and try again.")
        return

    # Initialize the Google Cloud Translation client
    translate_client = translate.Client()

    prs = Presentation(file_path)

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                original_text = shape.text
                if original_text.strip():
                    translated_text = translate_text(translate_client, original_text)
                    shape.text = translated_text
                    print(f"Translated: {original_text} -> {translated_text}")

    base, ext = os.path.splitext(file_path)
    output_file = base + "_translated" + ext
    prs.save(output_file)
    print(f"Translated file saved as: {output_file}")

if __name__ == "__main__":
    main()
